"""
Service d'import Excel pour les engagements hors-bilan (OffBalanceSheetItem).

Format attendu : un fichier Excel avec une feuille contenant les colonnes
listées dans EXPECTED_COLUMNS (détection par en-tête, case-insensitive).

L'import est :
- atomique par date d'arrêté,
- option replace=True pour supprimer les lignes du même date_arrete avant import.
- type_engagement invalide → forcé à "autre" avec warning dans le rapport.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from django.db import transaction

from .models import OffBalanceSheetItem, TYPE_CHOICES

VALID_TYPE_ENGAGEMENT = {c[0] for c in TYPE_CHOICES}

VALID_TYPE_TAUX = {"fixe", "variable", "administre", "indexe", "revisable"}

EXPECTED_COLUMNS = [
    "date_arrete", "type_engagement", "reference", "description",
    "client_id", "client_name", "segment", "secteur", "business_unit", "entite",
    "devise", "notionnel", "montant_utilise", "prob_tirage_pct",
    "taux", "type_taux", "date_mise_place", "maturite",
    "source",
]

HELPER_ROW_MARKERS = {
    "date d'arrêté",
    "yyyy-mm-dd",
    "ligne de crédit confirmée",
    "entreprise abc",
    "eng-2024-001",
}


@dataclass
class SheetReport:
    kind: str
    inserted: int = 0
    skipped: int = 0
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    sheets: list[SheetReport] = field(default_factory=list)

    @property
    def total_inserted(self) -> int:
        return sum(s.inserted for s in self.sheets)

    @property
    def total_errors(self) -> int:
        return sum(len(s.errors) for s in self.sheets)


def _parse_date(value: Any) -> dt.date | None:
    """Convertit une valeur Excel en date Python."""
    if value is None or value == "":
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
        raise ValueError(f"Date invalide : {value!r}")
    raise ValueError(f"Type de date non reconnu : {type(value)}")


def _is_template_helper_row(row: tuple[Any, ...]) -> bool:
    values = [str(value).strip().lower() for value in row if value not in (None, "")]
    if not values:
        return True
    joined = " | ".join(values)
    return any(marker in joined for marker in HELPER_ROW_MARKERS)


def import_off_balance_excel(file_obj, *, replace: bool = False) -> ImportReport:
    """
    Importe un fichier Excel contenant des engagements hors-bilan.

    :param file_obj: handle binaire (UploadedFile) ou chemin.
    :param replace: si True, supprime les lignes du même date_arrete avant import.
    :returns: ImportReport avec le détail des lignes insérées/erreurs.
    """
    workbook = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)
    report = ImportReport()

    # Prendre la feuille "hors_bilan" ou "off_balance" ou la première disponible
    sheet_name = workbook.sheetnames[0] if workbook.sheetnames else None
    for candidate in ("hors_bilan", "off_balance", "hors-bilan", "offbalance"):
        if candidate in [s.lower() for s in workbook.sheetnames]:
            sheet_name = next(s for s in workbook.sheetnames if s.lower() == candidate)
            break

    if not sheet_name:
        workbook.close()
        r = SheetReport(kind="hors_bilan")
        r.errors.append("Fichier Excel vide ou sans feuille.")
        report.sheets.append(r)
        return report

    sheet = workbook[sheet_name]
    r = SheetReport(kind=sheet_name)

    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        r.errors.append("Feuille vide.")
        report.sheets.append(r)
        workbook.close()
        return report

    # Mapping index -> nom_colonne (case-insensitive)
    headers_lower = [str(h).strip().lower() if h is not None else "" for h in raw_headers]
    col_map: dict[str, int] = {}
    for idx, h in enumerate(headers_lower):
        if h in EXPECTED_COLUMNS:
            col_map[h] = idx

    if not col_map:
        r.errors.append("Aucune colonne reconnue. Vérifiez les en-têtes.")
        report.sheets.append(r)
        workbook.close()
        return report

    objects: list[OffBalanceSheetItem] = []
    dates_seen: set[dt.date] = set()

    for row_idx, row in enumerate(rows, start=2):
        if _is_template_helper_row(row):
            continue

        def get(col, _row=row):
            idx = col_map.get(col)
            return _row[idx] if idx is not None and idx < len(_row) else None

        warnings: list[str] = []

        try:
            # Champs obligatoires
            date_arrete = _parse_date(get("date_arrete"))
            if date_arrete is None:
                raise ValueError("date_arrete est obligatoire.")

            # type_engagement — invalide → "autre" + warning
            type_raw = str(get("type_engagement") or "autre").strip().lower()
            if type_raw not in VALID_TYPE_ENGAGEMENT:
                warnings.append(
                    f"Ligne {row_idx} : type_engagement {type_raw!r} invalide, forcé à 'autre'."
                )
                type_raw = "autre"

            # type_taux
            type_taux_raw = str(get("type_taux") or "fixe").strip().lower()
            if type_taux_raw not in VALID_TYPE_TAUX:
                type_taux_raw = "fixe"

            # Montants
            notionnel_raw = get("notionnel")
            notionnel = int(float(notionnel_raw)) if notionnel_raw not in (None, "") else 0

            montant_utilise_raw = get("montant_utilise")
            montant_utilise = int(float(montant_utilise_raw)) if montant_utilise_raw not in (None, "") else 0

            prob_raw = get("prob_tirage_pct")
            prob_tirage_pct = float(prob_raw) if prob_raw not in (None, "") else 100.0

            taux_raw = get("taux")
            taux = float(taux_raw) if taux_raw not in (None, "") else 0.0

            obj = OffBalanceSheetItem(
                date_arrete=date_arrete,
                type_engagement=type_raw,
                reference=str(get("reference") or "").strip(),
                description=str(get("description") or "").strip(),
                client_id=str(get("client_id") or "").strip(),
                client_name=str(get("client_name") or "").strip(),
                segment=str(get("segment") or "").strip(),
                secteur=str(get("secteur") or "").strip(),
                business_unit=str(get("business_unit") or "").strip(),
                entite=str(get("entite") or "").strip(),
                devise=str(get("devise") or "XAF").strip()[:3],
                notionnel=notionnel,
                montant_utilise=montant_utilise,
                prob_tirage_pct=prob_tirage_pct,
                taux=taux,
                type_taux=type_taux_raw,
                date_mise_place=_parse_date(get("date_mise_place")),
                maturite=_parse_date(get("maturite")),
                source=str(get("source") or "flexcube").strip(),
            )
            objects.append(obj)
            dates_seen.add(date_arrete)
            # Ajouter les warnings au rapport (sans incrémenter skipped)
            r.errors.extend(warnings)

        except Exception as e:  # noqa: BLE001
            r.errors.append(f"Ligne {row_idx} : {e}")
            r.skipped += 1

    with transaction.atomic():
        if replace and dates_seen:
            OffBalanceSheetItem.objects.filter(date_arrete__in=dates_seen).delete()
        if objects:
            OffBalanceSheetItem.objects.bulk_create(objects, batch_size=500)
            r.inserted = len(objects)

    workbook.close()
    report.sheets.append(r)
    return report
