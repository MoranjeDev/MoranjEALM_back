"""
Service d'import Excel pour ClientMapping.

Colonnes attendues (case-insensitive, feuille "mapping" ou première feuille) :
  client_id (obligatoire), client_name, business_unit, segment, sous_segment,
  secteur, secteur_code, agence, entite.

Mode upsert : si client_id existe → update, sinon → create.
Retourne un ImportReport avec inserted, updated, errors.
"""
from __future__ import annotations

from dataclasses import dataclass, field

import openpyxl
from django.db import transaction

from .models import ClientMapping, SEGMENT_CHOICES, SOUS_SEGMENT_CHOICES

EXPECTED_COLUMNS = [
    "client_id", "client_name", "business_unit", "segment", "sous_segment",
    "secteur", "secteur_code", "agence", "entite",
]

VALID_SEGMENTS = {c[0] for c in SEGMENT_CHOICES}
VALID_SOUS_SEGMENTS = {c[0] for c in SOUS_SEGMENT_CHOICES}


@dataclass
class ImportReport:
    inserted: int = 0
    updated: int = 0
    errors: list[str] = field(default_factory=list)


def import_client_mapping_excel(file_obj) -> ImportReport:
    """
    Importe un fichier Excel de mapping clients MIS.

    :param file_obj: handle binaire (UploadedFile Django) ou chemin fichier.
    :returns: ImportReport avec inserted, updated, errors.
    """
    report = ImportReport()

    workbook = openpyxl.load_workbook(file_obj, data_only=True, read_only=True)

    # Chercher feuille "mapping" (case-insensitive) ou prendre la première
    sheet_name = workbook.sheetnames[0] if workbook.sheetnames else None
    for name in workbook.sheetnames:
        if name.strip().lower() == "mapping":
            sheet_name = name
            break

    if not sheet_name:
        workbook.close()
        report.errors.append("Fichier Excel vide ou sans feuille.")
        return report

    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)

    try:
        raw_headers = next(rows)
    except StopIteration:
        workbook.close()
        report.errors.append("Feuille vide.")
        return report

    # Mapping index → nom_colonne (case-insensitive)
    headers_lower = [str(h).strip().lower() if h is not None else "" for h in raw_headers]
    col_map: dict[str, int] = {}
    for idx, h in enumerate(headers_lower):
        if h in EXPECTED_COLUMNS:
            col_map[h] = idx

    if "client_id" not in col_map:
        workbook.close()
        report.errors.append("Colonne 'client_id' introuvable. Vérifiez les en-têtes.")
        return report

    to_insert: list[ClientMapping] = []
    to_update: list[tuple[ClientMapping, dict]] = []

    # Charger les client_id existants en mémoire pour éviter N+1
    existing: dict[str, ClientMapping] = {
        obj.client_id: obj
        for obj in ClientMapping.objects.all()
    }

    for row_idx, row in enumerate(rows, start=2):
        if not any(row):
            continue

        def get(col: str):
            idx = col_map.get(col)
            return row[idx] if idx is not None and idx < len(row) else None

        try:
            raw_id = get("client_id")
            if raw_id is None or str(raw_id).strip() == "":
                raise ValueError("client_id est obligatoire.")
            client_id = str(raw_id).strip()

            segment_raw = str(get("segment") or "other").strip().lower()
            if segment_raw not in VALID_SEGMENTS:
                segment_raw = "other"

            sous_segment_raw = str(get("sous_segment") or "").strip().lower()
            if sous_segment_raw not in VALID_SOUS_SEGMENTS:
                sous_segment_raw = ""

            fields = {
                "client_name":    str(get("client_name") or "").strip(),
                "business_unit":  str(get("business_unit") or "").strip(),
                "segment":        segment_raw,
                "sous_segment":   sous_segment_raw,
                "secteur":        str(get("secteur") or "").strip(),
                "secteur_code":   str(get("secteur_code") or "").strip(),
                "agence":         str(get("agence") or "").strip(),
                "entite":         str(get("entite") or "").strip(),
                "source":         "import_excel",
            }

            if client_id in existing:
                to_update.append((existing[client_id], fields))
            else:
                obj = ClientMapping(client_id=client_id, **fields)
                to_insert.append(obj)
                # Ajouter au dict pour dédupliquer dans le même fichier
                existing[client_id] = obj

        except Exception as exc:  # noqa: BLE001
            report.errors.append(f"Ligne {row_idx} : {exc}")

    workbook.close()

    with transaction.atomic():
        if to_insert:
            ClientMapping.objects.bulk_create(to_insert, batch_size=500)
            report.inserted = len(to_insert)

        update_fields = [
            "client_name", "business_unit", "segment", "sous_segment",
            "secteur", "secteur_code", "agence", "entite", "source",
        ]
        for obj, fields in to_update:
            for attr, val in fields.items():
                setattr(obj, attr, val)
        if to_update:
            ClientMapping.objects.bulk_update(
                [obj for obj, _ in to_update],
                update_fields,
                batch_size=500,
            )
            report.updated = len(to_update)

    return report
